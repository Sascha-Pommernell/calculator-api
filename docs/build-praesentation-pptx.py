#!/usr/bin/env python
"""Erzeugt docs/Calculator-API-Praesentation.pptx (Inhalt identisch zu praesentation.html).

    pip install python-pptx
    python docs/build-praesentation-pptx.py

Alle Folien bestehen aus nativen PowerPoint-Objekten (Textfelder, Tabellen,
Formen) und sind damit frei editierbar. Die Sprechernotizen landen in den
Notizenseiten der jeweiligen Folie.
"""

from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# ── Tokens ────────────────────────────────────────────────────────────────────
INK        = RGBColor(0x0B, 0x0B, 0x0B)   # primäre Schrift
INK_2      = RGBColor(0x52, 0x51, 0x4E)   # sekundäre Schrift
MUTED      = RGBColor(0x89, 0x87, 0x81)   # Labels, Fußzeile
HAIRLINE   = RGBColor(0xE1, 0xE0, 0xD9)   # Linien
SURFACE    = RGBColor(0xFC, 0xFC, 0xFB)   # Folienfläche
PLANE      = RGBColor(0xF9, 0xF9, 0xF7)   # Kartenfläche
CODE_BG    = RGBColor(0xF4, 0xF4, 0xF1)
BLUE       = RGBColor(0x2A, 0x78, 0xD6)   # Akzent / Balken (Slot 1)
AQUA       = RGBColor(0x1B, 0xAF, 0x7A)   # Strings im Code (Slot 3)
VIOLET     = RGBColor(0x4A, 0x3A, 0xA7)   # Typen / Zahlen im Code (Slot 7)
GOOD       = RGBColor(0x0C, 0xA3, 0x0C)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)

SANS = "Segoe UI"
MONO = "Consolas"

# ── Geometrie (16:9) ─────────────────────────────────────────────────────────
SW, SH   = 13.333, 7.5
L        = 0.80                 # linker Rand
CW       = SW - 2 * L           # Inhaltsbreite
Y_KICKER = 0.52
Y_TITLE  = 0.80
Y_RULE   = 1.46
Y_BODY   = 1.74
Y_FOOT   = 6.94
BODY_H   = Y_FOOT - Y_BODY - 0.12

GAP   = 0.30
COL_2 = (CW - GAP) / 2
COL_3 = (CW - 2 * 0.24) / 3

prs = Presentation()
prs.slide_width  = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]


# ── Bausteine ─────────────────────────────────────────────────────────────────
def slide(notes: str = "") -> "object":
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SW), Inches(SH))
    bg.fill.solid()
    bg.fill.fore_color.rgb = SURFACE
    bg.line.fill.background()
    bg.shadow.inherit = False
    if notes:
        s.notes_slide.notes_text_frame.text = notes.strip()
    return s


def box(s, x, y, w, h, *, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, *, first=False, space_before=0, space_after=0, line=1.28,
         align=PP_ALIGN.LEFT, bullet_indent=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line
    p.alignment = align
    if bullet_indent is not None:
        p.level = 0
    return p


def run(p, text, *, size=11, color=INK_2, bold=False, italic=False, font=SANS):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font
    return r


def rich(p, text, *, size=11, color=INK_2, font=SANS):
    """**fett** und `code` in einem Fließtext auszeichnen (auch verschachtelt)."""
    for part in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            # `code` innerhalb von **fett** behält beides
            for sub in re.split(r"(`[^`]+`)", part[2:-2]):
                if not sub:
                    continue
                if sub.startswith("`"):
                    run(p, sub[1:-1], size=size * 0.94, color=INK,
                        bold=True, font=MONO)
                else:
                    run(p, sub, size=size, color=INK, bold=True, font=font)
        elif part.startswith("`"):
            run(p, part[1:-1], size=size * 0.94, color=INK, font=MONO)
        else:
            run(p, part, size=size, color=color, font=font)


def bullet_char(p, *, char="▪", color=BLUE, indent=0.26):
    """Native PowerPoint-Aufzählung mit hängendem Einzug.

    Damit übernimmt PowerPoint den Zeilenumbruch und die Absatzabstände –
    keine geschätzten Zeilenhöhen, kein Überlappen bei langen Texten.
    """
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent))))
    pPr.set("indent", str(int(-Inches(indent))))
    buClr = pPr.makeelement(qn("a:buClr"), {})
    buClr.append(pPr.makeelement(qn("a:srgbClr"), {"val": f"{color}"}))
    pPr.append(buClr)
    pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": char}))


def rect(s, x, y, w, h, *, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE,
         line_w=0.75):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = 0.06
    sp.text_frame.word_wrap = True
    return sp


def hline(s, x, y, w, color=HAIRLINE):
    rect(s, x, y, w, 0.011, fill=color)


def head(s, kicker, title, *, foot=None, pageno=None):
    tf = box(s, L, Y_KICKER, CW, 0.26)
    run(para(tf, first=True), kicker.upper(), size=10.5, color=BLUE, bold=True)

    tf = box(s, L, Y_TITLE, CW, 0.62)
    run(para(tf, first=True, line=1.05), title, size=27, color=INK, bold=True)

    rect(s, L, Y_RULE, 0.55, 0.035, fill=BLUE)

    if foot:
        hline(s, L, Y_FOOT, CW)
        tf = box(s, L, Y_FOOT + 0.10, CW * 0.6, 0.22)
        run(para(tf, first=True), foot, size=9, color=MUTED)
        tf = box(s, L + CW * 0.6, Y_FOOT + 0.10, CW * 0.4, 0.22)
        run(para(tf, first=True, align=PP_ALIGN.RIGHT),
            f"{pageno} / 25", size=9, color=MUTED)


def bullets(s, x, y, w, items, *, size=11.5, gap=9, dot=BLUE, h=None):
    """items: Liste von Strings (mit **fett** / `code`) oder (text, [sub…])."""
    tf = box(s, x, y, w, h if h is not None else Y_FOOT - y - 0.15)
    first = True
    for item in items:
        text, subs = (item, []) if isinstance(item, str) else item
        p = para(tf, first=first, space_after=gap, line=1.26)
        first = False
        bullet_char(p, color=dot)
        rich(p, text, size=size)
        for sub in subs:
            sp = para(tf, space_after=4, line=1.22)
            bullet_char(sp, char="–", color=MUTED, indent=0.44)
            run(sp, sub, size=size - 1.5, color=MUTED)
    return tf


def card(s, x, y, w, h, title, body, *, accent=None, body_size=10.5):
    rect(s, x, y, w, h, fill=PLANE, line=HAIRLINE,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    if accent:
        rect(s, x, y + 0.06, 0.035, h - 0.12, fill=accent)
    pad = 0.22
    cy = y + 0.17
    if title:
        tf = box(s, x + pad, cy, w - 2 * pad, 0.28)
        run(para(tf, first=True), title, size=12, color=INK, bold=True)
        cy += 0.34
    tf = box(s, x + pad, cy, w - 2 * pad, h - (cy - y) - 0.16)
    for i, ptext in enumerate(body if isinstance(body, list) else [body]):
        rich(para(tf, first=(i == 0), space_before=0 if i == 0 else 5, line=1.3),
             ptext, size=body_size)


# ── Code-Blöcke ───────────────────────────────────────────────────────────────
KEYWORDS = {
    "public", "private", "static", "return", "if", "throw", "new", "string",
    "double", "void", "class", "nameof", "using", "bool", "var", "sealed",
    "record", "required", "init", "get", "namespace", "true", "false", "null",
    "FROM", "RUN", "COPY", "ENV", "EXPOSE", "USER", "HEALTHCHECK",
    "ENTRYPOINT", "CMD", "AS", "WORKDIR", "exit",
}
TYPES = {
    "ActionResult", "CalculationResponse", "CalculationRequest", "ProblemDetails",
    "HttpPost", "ProducesResponseType", "FromBody", "OperationNames", "Func",
    "IReadOnlyList", "DivideByZeroException", "OverflowException", "Theory",
    "Fact", "InlineData", "MemberData", "Assert", "ArgumentException",
    "Status200OK", "Status400BadRequest", "CalculatorService", "Calculate",
}
TOKEN_RE = re.compile(r'("(?:[^"\\]|\\.)*"|\b[A-Za-z_][A-Za-z0-9_]*\b|\b\d+(?:\.\d+)?(?:e-?\d+)?\b)')


CODE_PAD = 0.30


def code_h(lines, size=9.5):
    """Höhe eines Code-Blocks aus der Zeilenzahl.

    PowerPoint multipliziert den Zeilenabstand mit der Fontzeilenhöhe (≈1.2 em),
    daher der Faktor 1.32 × 1.2 – so passt der Rahmen immer zum Inhalt.
    """
    return len(lines) * size * 1.32 * 1.2 / 72 + CODE_PAD


def code(s, x, y, w, lines, *, size=9.5, label=None):
    h = code_h(lines, size)
    if label:
        tf = box(s, x, y - 0.26, w, 0.22)
        run(para(tf, first=True), label, size=9, color=MUTED, font=MONO)
    rect(s, x, y, w, h, fill=CODE_BG, line=HAIRLINE,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = box(s, x + 0.20, y + 0.15, w - 0.34, h - 0.30, wrap=False)
    for i, line in enumerate(lines):
        p = para(tf, first=(i == 0), line=1.32)
        stripped = line.lstrip(" ")
        if stripped.startswith(("//", "#")):
            run(p, line, size=size, color=MUTED, italic=True, font=MONO)
            continue
        # Kommentar am Zeilenende abtrennen
        head_txt, comment = line, ""
        m = re.search(r"(?<!:)//", line)
        if m and line.count('"') % 2 == 0:
            head_txt, comment = line[:m.start()], line[m.start():]
        pos = 0
        for m in TOKEN_RE.finditer(head_txt):
            if m.start() > pos:
                run(p, head_txt[pos:m.start()], size=size, color=INK, font=MONO)
            tok = m.group(0)
            if tok.startswith('"'):
                col = AQUA
            elif tok in KEYWORDS:
                col = BLUE
            elif tok in TYPES:
                col = VIOLET
            elif tok[0].isdigit():
                col = VIOLET
            else:
                col = INK
            run(p, tok, size=size, color=col, font=MONO)
            pos = m.end()
        if pos < len(head_txt):
            run(p, head_txt[pos:], size=size, color=INK, font=MONO)
        if comment:
            run(p, comment, size=size, color=MUTED, italic=True, font=MONO)


# ── Tabelle ───────────────────────────────────────────────────────────────────
NO_GRID_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # „Kein Stil, kein Gitter“


def _strip_table_style(tbl):
    """Standardgitter entfernen – Trennlinien werden explizit gesetzt."""
    tblPr = tbl._tbl.tblPr
    for el in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(el)
    el = tblPr.makeelement(qn("a:tableStyleId"), {})
    el.text = NO_GRID_STYLE
    tblPr.append(el)


def _rule_below(cell, color=HAIRLINE):
    """Einzelne Unterlinie an einer Zelle – entspricht border-bottom im HTML."""
    tcPr = cell._tc.get_or_add_tcPr()
    ln = tcPr.makeelement(qn("a:lnB"), {"w": str(int(Pt(0.75))), "cap": "flat",
                                        "cmpd": "sng", "algn": "ctr"})
    fill = ln.makeelement(qn("a:solidFill"), {})
    fill.append(ln.makeelement(qn("a:srgbClr"), {"val": f"{color}"}))
    ln.append(fill)
    tcPr.insert(0, ln)   # lnB muss laut Schema vor den Füll-Elementen stehen


def table(s, x, y, w, widths, rows, *, header, row_h=0.38, size=10.5):
    n = len(rows) + 1
    shp = s.shapes.add_table(n, len(widths), Inches(x), Inches(y),
                             Inches(w), Inches(row_h * n))
    tbl = shp.table
    tbl.first_row = False
    tbl.horz_banding = False
    _strip_table_style(tbl)
    for i, frac in enumerate(widths):
        tbl.columns[i].width = Emu(int(Inches(w) * frac))

    for c, text in enumerate(header):
        cell = tbl.cell(0, c)
        cell.margin_left = Inches(0); cell.margin_right = Inches(0.18)
        cell.margin_top = Inches(0.02); cell.margin_bottom = Inches(0.10)
        _rule_below(cell)
        cell.fill.background()
        run(para(cell.text_frame, first=True), text.upper(),
            size=9, color=MUTED, bold=True)
    tbl.rows[0].height = Inches(0.30)

    for r, cells in enumerate(rows, start=1):
        tbl.rows[r].height = Inches(row_h)
        for c, text in enumerate(cells):
            cell = tbl.cell(r, c)
            cell.margin_left = Inches(0); cell.margin_right = Inches(0.18)
            cell.margin_top = Inches(0.08); cell.margin_bottom = Inches(0.08)
            cell.fill.background()
            cell.vertical_anchor = MSO_ANCHOR.TOP
            rich(para(cell.text_frame, first=True, line=1.22), text, size=size)
    return tbl


# ═════════════════════════════════════════════════════════════════════════════
#  Folien
# ═════════════════════════════════════════════════════════════════════════════

# 1 — Titel
s = slide("""
Rahmung in einem Satz: „Ich zeige heute eine REST-API für Grundrechenarten. Das
Spannende daran ist nicht das Rechnen – sondern alles drumherum: wie ich
sicherstelle, dass sie korrekt ist, und wie das automatisiert bei jedem Push
überprüft wird.“

Zeitbudget bei 20 Minuten: Teil 1 (Folien 1–12) ca. 9 Min., Teil 2 (13–21)
ca. 8 Min., Abschluss (22–25) ca. 3 Min. Bei 10 Minuten die Folien 5, 8, 15, 19
und 24 überspringen – der rote Faden bleibt erhalten.
""")
tf = box(s, 1.0, 1.62, 10, 0.3)
run(para(tf, first=True), "PROJEKTPRÄSENTATION", size=11, color=BLUE, bold=True)
tf = box(s, 1.0, 2.02, 10, 0.95)
run(para(tf, first=True, line=1.0), "Calculator API", size=50, color=INK, bold=True)
tf = box(s, 1.0, 3.12, 8.6, 1.1)
rich(para(tf, first=True, line=1.3),
     "Eine REST-API für Grundrechenarten – und ein vollständiger "
     "Qualitätssicherungsprozess von der Testkonzeption bis zum "
     "veröffentlichten Testbericht.", size=15)
bx = 1.0
for tag in ["ASP.NET Core · .NET 10", "xUnit", "NUnit + Playwright", "Docker",
            "GitHub Actions", "Allure"]:
    w = 0.070 * len(tag) + 0.40
    sp = rect(s, bx, 4.48, w, 0.36, fill=RGBColor(0xCD, 0xE2, 0xFB),
              shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    sp.adjustments[0] = 0.5
    tfb = sp.text_frame
    tfb.word_wrap = False
    tfb.margin_left = tfb.margin_right = Inches(0.06)
    tfb.vertical_anchor = MSO_ANCHOR.MIDDLE
    run(para(tfb, first=True, align=PP_ALIGN.CENTER), tag,
        size=9.5, color=RGBColor(0x0D, 0x36, 0x6B), bold=True)
    bx += w + 0.13
tf = box(s, 1.0, 5.35, 8, 0.7)
run(para(tf, first=True, line=1.5), "Sascha Pommernell", size=11, color=MUTED)
run(para(tf, line=1.5), "github.com/Sascha-Pommernell/calculator-api",
    size=11, color=MUTED)

# 2 — Agenda
s = slide("""
Nur überfliegen. Sagen, dass die Präsentation zwei Teile hat: die API und
wie sie abgesichert wird.
""")
head(s, "Überblick", "Agenda", foot="Calculator API", pageno=2)
bullets(s, L, Y_BODY, COL_2, [
    "**Ziel des Projekts** – warum ein Taschenrechner?",
    "**Architektur & Tech-Stack**",
    "**Die API** – Endpunkte, Verträge, Code",
    "**Robustheit** – Fehlerbehandlung & Härtung",
], size=13, gap=16)
bullets(s, L + COL_2 + GAP, Y_BODY, COL_2, [
    "**Qualitätssicherung** – Testkonzept nach ISTQB",
    "**Testautomatisierung** – Unit- und API-Tests",
    "**CI/CD** – Docker & GitHub Actions",
    "**Ergebnis & Ausblick**",
], size=13, gap=16)

# 3 — Motivation
s = slide("""
Wichtigste Folie für die Erwartungssteuerung. Hier vorwegnehmen, was das
Publikum sonst denkt („ein Taschenrechner, ernsthaft?“):

„Die Fachlichkeit ist absichtlich trivial. Wenn ich hier ein komplexes
Domänenmodell hätte, würde die halbe Zeit für Erklärungen draufgehen. So kann ich
stattdessen zeigen, wie ein Projekt aufgebaut ist, das man auch in zwei Jahren
noch anfassen kann.“
""")
head(s, "Motivation", "Einfache Fachlichkeit, professioneller Prozess",
     foot="Ziel des Projekts", pageno=3)
tf = box(s, L, Y_BODY, CW - 1.2, 0.9)
rich(para(tf, first=True, line=1.34),
     "Die Rechenlogik ist absichtlich **trivial** – dadurch liegt der Fokus "
     "vollständig auf dem, was in echten Projekten den Unterschied macht: "
     "Architektur, Verträge, Fehlerbehandlung, Testkonzeption, Automatisierung "
     "und Nachvollziehbarkeit.", size=13)
cy = Y_BODY + 1.12
for i, (t, b) in enumerate([
    ("Was gebaut wurde",
     "Eine REST-API für Addition, Subtraktion, Multiplikation und Division mit "
     "**zwei oder mehr** Zahlen."),
    ("Was daran interessant ist",
     "Jede Zeile Produktivcode ist durch ein **dokumentiertes Testkonzept** und "
     "eine automatisierte Pipeline abgesichert."),
    ("Was daran gelernt wurde",
     "Testentwurf nach ISTQB, Containerisierung, CI/CD und **Reporting** als "
     "durchgängige Kette."),
]):
    card(s, L + i * (COL_3 + 0.24), cy, COL_3, 1.85, t, b)

# 4 — Zwei Repositories
s = slide("""
Kernbotschaft: Trennung erzwingt Black-Box-Disziplin. Wären die Tests im gleichen
Projekt, könnte man versehentlich interne Klassen aufrufen und würde dann nicht
mehr die API testen, sondern den Code.

Wenn gefragt wird, ob das nicht unnötig kompliziert ist: Ja, für dieses Projekt
ist es etwas Overhead – aber genau so trennen echte Projekte Test- und
Produktivcode, wenn ein separates QA-Team existiert.
""")
head(s, "Systemüberblick", "Zwei Repositories, eine Pipeline",
     foot="Architektur", pageno=4)
card(s, L, Y_BODY, COL_2, 2.05, "calculator-api  ·  Produktivcode", [
    "ASP.NET-Core-Projekt (4 Endpunkte + Health)",
    "Unit-Tests `Calculator.Api.Tests` (xUnit)",
    "Dockerfile · CI-Workflow · Testkonzept",
], accent=BLUE)
card(s, L, Y_BODY + 2.25, COL_2, 2.05,
     "calculator-api-tests-c  ·  Testautomatisierung", [
    "API-Tests mit NUnit 4 + Playwright",
    "Allure- und TRX-Reporting",
    "Eigenes Docker-Image, eigene Pipeline",
], accent=AQUA)
x2 = L + COL_2 + GAP
tf = box(s, x2, Y_BODY, COL_2, 0.25)
run(para(tf, first=True), "Warum getrennt?", size=9, color=MUTED, font=MONO)
bullets(s, x2, Y_BODY + 0.36, COL_2, [
    "**Black-Box-Disziplin.** Die Tests kennen nur den HTTP-Vertrag – keine "
    "internen Klassen, kein Zugriff auf die Implementierung.",
    "**Unabhängige Lebenszyklen.** Testcode kann weiterentwickelt werden, ohne "
    "die API zu berühren – und umgekehrt.",
    "**Beidseitiges Gate.** Jedes Repo hat einen Workflow, der jeweils das andere "
    "auscheckt: Änderungen an der API und am Testcode werden abgesichert.",
], size=11.5, gap=15)

# 5 — Tech-Stack
s = slide("""
Nicht vorlesen. Nur zwei Punkte hervorheben:

- Playwright ohne Browser – überrascht viele. Playwright kann auch reines HTTP
  (APIRequestContext). Man bekommt Parallelisierung und Tracing gratis.
- Allure – der Grund, warum am Ende ein Bericht existiert, den auch
  Nichtentwickler lesen können.
""")
head(s, "Technologie", "Tech-Stack", foot="Architektur", pageno=5)
table(s, L, Y_BODY, CW, [0.22, 0.28, 0.50],
      header=["Bereich", "Technologie", "Begründung"],
      rows=[
        ["Laufzeit / Framework", "**.NET 10 · ASP.NET Core**",
         "Aktuelle LTS-Generation, minimaler Boilerplate, integriertes DI"],
        ["API-Dokumentation", "**OpenAPI + Swagger UI**",
         "Vertrag ist maschinenlesbar und im Browser ausprobierbar"],
        ["Unit-Tests", "**xUnit**",
         "Datengetriebene `[Theory]`-Tests für die Rechenlogik"],
        ["API-Tests", "**NUnit 4 + Playwright**",
         "`APIRequestContext` – HTTP-Tests ohne Browser, parallelisierbar"],
        ["Reporting", "**Allure + TRX**",
         "Visueller Bericht mit Historie, dazu ein Check am Commit"],
        ["Betrieb", "**Docker** (Multi-Stage)",
         "Identische Umgebung lokal und in der CI"],
        ["CI/CD", "**GitHub Actions**",
         "Build, Test und Veröffentlichung bei jedem Push"],
      ], row_h=0.55, size=11)

# 6 — Schichten
s = slide("""
Am Baum entlanggehen. Die Pointe steht in der Karte unten rechts: weil die
Fachlogik nicht am HTTP-Stack hängt, laufen 33 Tests in 46 Millisekunden. Das ist
der praktische Nutzen der Schichtentrennung – nicht die Theorie.
""")
head(s, "Aufbau", "Schichten & Verantwortlichkeiten", foot="Architektur", pageno=6)
code(s, L, Y_BODY, 6.15, [
    "Calculator.Api/",
    "├── Controllers/",
    "│   └── CalculatorController.cs   // HTTP-Schicht",
    "├── Services/",
    "│   ├── ICalculatorService.cs     // Vertrag",
    "│   └── CalculatorService.cs      // Fachlogik",
    "├── Models/",
    "│   ├── CalculationRequest.cs     // + Validierung",
    "│   ├── CalculationResponse.cs",
    "│   └── OperationNames.cs",
    "├── Infrastructure/",
    "│   └── CalculationExceptionHandler.cs",
    "└── Program.cs                    // Komposition",
], size=10.5)
x2 = L + 6.45
bullets(s, x2, Y_BODY, CW - 6.45, [
    "**Controller** nimmt HTTP an und gibt HTTP zurück – keine Rechenlogik, "
    "keine `try/catch`-Blöcke.",
    "**Service** enthält die vollständige Fachlogik und ist damit isoliert und "
    "ohne Webserver testbar.",
    "**Interface + Dependency Injection** – der Controller hängt an "
    "`ICalculatorService`, nicht an der Implementierung.",
    "**Infrastructure** übersetzt Ausnahmen zentral in einheitliche "
    "Fehlerantworten.",
], size=11.5, gap=13)
card(s, x2, Y_BODY + 3.05, CW - 6.45, 0.85, None,
     "**Konsequenz:** Die Fachlogik wird von 33 Unit-Tests in Millisekunden "
     "geprüft – ganz ohne HTTP.", body_size=11)

# 7 — Endpunkte
s = slide("""
Die interessante Stelle ist die Entwurfsentscheidung unten links: vier Endpunkte
statt /calculate?op=add. Begründung: der Vertrag wird selbstdokumentierend, jeder
Endpunkt hat eigene OpenAPI-Beschreibung, und es gibt keinen String, den man
falsch schreiben kann.

Zweite Pointe: Es sind Listen, keine zwei Operanden – [1,2,3,4] geht.
""")
head(s, "Die API", "Endpunkte", foot="Die API", pageno=7)
table(s, L, Y_BODY, CW, [0.32, 0.11, 0.57],
      header=["Endpunkt", "Methode", "Verhalten"],
      rows=[
        ["`/api/calculator/add`", "POST", "Addiert alle übergebenen Zahlen"],
        ["`/api/calculator/subtract`", "POST",
         "Subtrahiert alle weiteren Zahlen von der ersten (links-assoziativ)"],
        ["`/api/calculator/multiply`", "POST",
         "Multipliziert alle übergebenen Zahlen"],
        ["`/api/calculator/divide`", "POST",
         "Dividiert die erste Zahl nacheinander durch alle weiteren"],
        ["`/health`", "GET",
         "Erreichbarkeits-Probe für Container, CI und Tests"],
      ], row_h=0.40, size=11)
cy = Y_BODY + 2.55
card(s, L, cy, COL_2, 1.55, "Bewusste Entwurfsentscheidung",
     "Ein Endpunkt **pro Operation** statt eines generischen "
     "`/calculate?op=…`: der Vertrag ist selbsterklärend, in OpenAPI einzeln "
     "dokumentiert und ohne Magic Strings testbar.")
card(s, L + COL_2 + GAP, cy, COL_2, 1.55, "Beliebig viele Operanden",
     "Alle Endpunkte arbeiten auf einer **Liste** – nicht auf `a` und `b`. "
     "`[1,2,3,4]` ist genauso gültig wie `[1,2]`.")

# 8 — Request / Response
s = slide("""
Nur kurz. Erklären, warum die Antwort die Eingabe zurückspiegelt: bei asynchronen
oder parallelen Aufrufen kann der Client Anfrage und Ergebnis zuordnen, ohne
selbst mitzuzählen.
""")
head(s, "Die API", "Der Vertrag: Request & Response", foot="Die API", pageno=8)
code(s, L, Y_BODY + 0.28, COL_2,
     ['{', '  "numbers": [10, 5, 2]', '}'],
     size=12, label="POST /api/calculator/divide")
tf = box(s, L, Y_BODY + 1.72, COL_2, 0.7)
rich(para(tf, first=True, line=1.32),
     "Mindestens **2**, höchstens **1000** Zahlen. "
     "`Content-Type: application/json`.", size=12)
x2 = L + COL_2 + GAP
code(s, x2, Y_BODY + 0.28, COL_2,
     ['{', '  "operation": "Division",', '  "numbers": [10, 5, 2],',
      '  "result": 1', '}'],
     size=12, label="200 OK")
tf = box(s, x2, Y_BODY + 2.15, COL_2, 1.0)
rich(para(tf, first=True, line=1.32),
     "Die Antwort **spiegelt die Eingabe zurück** – der Aufrufer kann Anfrage "
     "und Ergebnis eindeutig zuordnen, auch bei parallelen Requests.", size=12)

# 9 — Controller
s = slide("""
Hier langsamer werden – erste Code-Folie. Zwei Dinge betonen:

1. Eine Hilfsmethode für alle vier Operationen, die Operation kommt als Delegate
   herein. Kein Copy-Paste x4.
2. Kein try/catch. Bewusst. Das führt zur nächsten Folie über zentrale
   Fehlerbehandlung.
""")
head(s, "Code", "Der Controller bleibt dünn", foot="Die API", pageno=9)
code(s, L, Y_BODY + 0.28, 6.5, [
    '[HttpPost("divide")]',
    '[ProducesResponseType<CalculationResponse>(Status200OK)]',
    '[ProducesResponseType<ProblemDetails>(Status400BadRequest)]',
    'public ActionResult<CalculationResponse> Divide(',
    '        [FromBody] CalculationRequest request)',
    '    => Calculate(OperationNames.Division,',
    '                 request,',
    '                 _calculatorService.Divide);',
    '',
    '// Fachliche Ausnahmen behandelt zentral der',
    '// CalculationExceptionHandler.',
    'private ActionResult<CalculationResponse> Calculate(',
    '        string operation,',
    '        CalculationRequest request,',
    '        Func<IReadOnlyList<double>, double> calculation)',
    '    => Ok(new CalculationResponse(',
    '           operation, request.Numbers,',
    '           calculation(request.Numbers)));',
], size=9.5, label="Controllers/CalculatorController.cs")
x2 = L + 6.8
bullets(s, x2, Y_BODY + 0.28, CW - 6.8, [
    "**Eine gemeinsame Hilfsmethode** für alle vier Operationen – die Operation "
    "wird als Delegate übergeben. Kein vierfach kopierter Code.",
    "**Kein `try/catch`.** Fehler werden bewusst nach oben durchgelassen und an "
    "genau einer Stelle behandelt.",
    "**`ProducesResponseType`** macht die möglichen Antworten Teil des "
    "OpenAPI-Dokuments – der Vertrag steht im Code, nicht in einer separaten Datei.",
    "**Kein Zustand.** Der Service ist als Singleton registriert und vollständig "
    "zustandslos – dadurch sind parallele Requests und parallele Tests unkritisch.",
], size=11.5, gap=15)

# 10 — Service
s = slide("""
Die inhaltlich stärkste Code-Folie. Zwei Fallen im Umgang mit double:

„In C# wirft die Division 10.0 / 0 keine Exception – sie liefert Infinity. Wer
sich auf DivideByZeroException verlässt, hat einen Bug. Also prüfe ich den Divisor
selbst, vor der Rechnung.“

„Und selbst wenn kein Divisor null ist, kann das Ergebnis aus dem darstellbaren
Bereich laufen. EnsureFinite fängt das ab – sonst würde die API Infinity
ausliefern, was in JSON überhaupt kein gültiger Wert ist.“

Hier ankündigen: „Diese Prüfung ist nicht aus Weitsicht entstanden, sondern weil
ein Test fehlgeschlagen ist – dazu komme ich auf Folie 22.“
""")
head(s, "Code", "Die Fachlogik im Service", foot="Die API", pageno=10)
code(s, L, Y_BODY + 0.28, 6.5, [
    'public double Divide(IReadOnlyList<double> numbers)',
    '{',
    '    ValidateInput(numbers);',
    '',
    '    if (numbers.Skip(1).Any(n => n == 0))',
    '    {',
    '        throw new DivideByZeroException(',
    '            "Division durch null ist nicht erlaubt.");',
    '    }',
    '',
    '    return EnsureFinite(numbers.Skip(1)',
    '        .Aggregate(numbers[0], (r, n) => r / n));',
    '}',
    '',
    'private static double EnsureFinite(double result)',
    '{',
    '    if (!double.IsFinite(result))',
    '    {',
    '        throw new OverflowException(…);',
    '    }',
    '    return result;',
    '}',
], size=9.0, label="Services/CalculatorService.cs")
x2 = L + 6.8
bullets(s, x2, Y_BODY + 0.28, CW - 6.8, [
    "**Divisor-Prüfung vor der Rechnung.** C# liefert bei `double`-Division "
    "durch 0 kein Exception, sondern `Infinity` – der Fehler muss aktiv erkannt "
    "werden.",
    "**`EnsureFinite` als Torwächter.** Jede Operation prüft ihr Ergebnis auf "
    "`IsFinite`. Damit kann die API nie `Infinity` oder `NaN` ausliefern – Werte, "
    "die in JSON ohnehin nicht gültig sind.",
    "**Herkunft dieser Prüfung:** Sie entstand als Reaktion auf einen Testfall, "
    "der ursprünglich fehlschlug – siehe Folie „Herausforderungen“.",
    "**Fehler als Ausnahmen** statt als Rückgabewert-Codes: der Aufrufer kann sie "
    "nicht versehentlich ignorieren.",
], size=11.5, gap=15)

# 11 — ProblemDetails
s = slide("""
Kernbotschaft: einheitliches Fehlerformat, standardisiert nach RFC 9457.

Der Unterschied zwischen den beiden Beispielen ist wichtig:
- links: Validierung greift, bevor der Code läuft → errors-Objekt pro Feld
- rechts: fachlicher Fehler im Service → title + detail

Abschließender Satz: „Der Client braucht keine Textanalyse. Und weil die Struktur
Teil des Vertrags ist, wird sie getestet – nicht nur dokumentiert.“
""")
head(s, "Robustheit", "Einheitliche Fehlerantworten nach RFC 9457",
     foot="Robustheit", pageno=11)
tf = box(s, L, Y_BODY, CW - 0.9, 0.75)
rich(para(tf, first=True, line=1.32),
     "Ein zentraler `IExceptionHandler` übersetzt fachliche Ausnahmen in "
     "**ProblemDetails**. Jede Fehlerantwort der API – auch 404, 405, 415 und "
     "unerwartete 500er – hat damit dasselbe, maschinenlesbare Format.", size=12)
cy = Y_BODY + 1.15
code(s, L, cy, COL_2, [
    '{',
    '  "title": "One or more validation',
    '            errors occurred.",',
    '  "status": 400,',
    '  "errors": {',
    '    "Numbers": ["Es müssen mindestens',
    '                zwei Zahlen …"]',
    '  }',
    '}',
], size=9.5, label="Eingabevalidierung → ValidationProblemDetails")
code(s, L + COL_2 + GAP, cy, COL_2, [
    '{',
    '  "title": "Ungültige Berechnung",',
    '  "status": 400,',
    '  "detail": "Division durch null ist',
    '             nicht erlaubt."',
    '}',
], size=9.5, label="Fachlicher Fehler → ProblemDetails")
card(s, L, cy + 2.40, CW, 0.9, None,
     "**Warum das wichtig ist:** Der Client braucht keine Fallunterscheidung "
     "nach Text. Die Struktur ist Teil des Vertrags – und wird von den "
     "Vertragstests geprüft, nicht nur dokumentiert.", body_size=11.5)

# 12 — Härtung
s = slide("""
Sechs Karten, jede in einem Satz. Nicht alle gleich lang behandeln – Schwerpunkt:

- Rate Limiting: die API würde ohne Limit von einer Schleife lahmgelegt.
- non-root im Container: eine Zeile USER app, aber deutlich kleinere
  Angriffsfläche, wenn der Prozess kompromittiert wird.
- Swagger nur in Development: in Produktion soll die API ihre eigene Landkarte
  nicht ausliefern.
""")
head(s, "Robustheit", "Härtung & Betriebsreife", foot="Robustheit", pageno=12)
cards = [
    ("Eingabegrenzen",
     "Deklarative Validierung am Model: `[MinLength(2)]` und `[MaxLength(1000)]`. "
     "Die Obergrenze verhindert, dass ein einzelner Request beliebig viel Arbeit "
     "erzeugt."),
    ("Zahlbereich",
     "`double.IsFinite` auf jedes Ergebnis. Arithmetischer Überlauf wird als "
     "**400** abgelehnt statt als `Infinity` ausgeliefert."),
    ("Rate Limiting",
     "**200 Requests/Sekunde pro Client-IP** über einen Fixed-Window-Limiter; "
     "darüber `429 Too Many Requests`."),
    ("Health-Endpoint",
     "`GET /health` dient dem Docker-`HEALTHCHECK`, der CI-Wartelogik und der "
     "Entry-Prüfung der Tests."),
    ("Container-Sicherheit",
     "Das Image läuft als **non-root-User** `app` – kein Root-Prozess im "
     "Container."),
    ("Swagger nur in Development",
     "OpenAPI-Dokument und Swagger UI sind an `IsDevelopment()` gebunden und in "
     "Produktion nicht exponiert."),
]
for i, (t, b) in enumerate(cards):
    card(s, L + (i % 3) * (COL_3 + 0.24), Y_BODY + (i // 3) * 2.28,
         COL_3, 2.05, t, b)

# 13 — Abschnittstrenner
s = slide("""
Atempause. Ein Satz: „Damit ist die API beschrieben. Jetzt der eigentliche Kern:
wie ich beweise, dass sie funktioniert.“
""")
tf = box(s, L, 2.55, 8, 0.3)
run(para(tf, first=True), "TEIL 2", size=11, color=BLUE, bold=True)
tf = box(s, L, 2.95, 10, 0.85)
run(para(tf, first=True, line=1.05), "Qualitätssicherung", size=40,
    color=INK, bold=True)
tf = box(s, L, 4.05, 8.2, 0.8)
rich(para(tf, first=True, line=1.32),
     "Vom dokumentierten Testkonzept über automatisierte Tests bis zum "
     "veröffentlichten Bericht.", size=15)

# 14 — Zwei Testebenen
s = slide("""
Die Testpyramide erklären: viele schnelle Tests unten, wenige teure oben.

Wichtigster Punkt – „kein grün ohne Tests“: „Ein früher Zustand meiner Pipeline
war grün, obwohl gar keine Tests gelaufen waren: die API war nicht erreichbar,
und die Tests haben sich selbst übersprungen. Das ist schlimmer als ein roter
Build, weil man ihm glaubt. Jetzt erzwingt CI=true einen harten Fehlschlag.“

Zweiter Punkt: In der CI wird gegen den Docker-Container in
Produktionskonfiguration getestet, nicht gegen dotnet run. Getestet wird das
Artefakt, das ausgeliefert würde.
""")
head(s, "Testkonzept", "Zwei Testebenen mit klarer Aufgabenteilung",
     foot="Qualitätssicherung", pageno=14)
card(s, L, Y_BODY, COL_2, 2.05, "API-Tests  ·  Black Box", [
    "**NUnit 4 + Playwright** gegen die laufende Anwendung im Container. "
    "Getestet wird ausschließlich über HTTP.",
    "Prüft: Statuscodes, Response-Struktur, Header, Routen, HTTP-Methoden – "
    "also den **Vertrag**.",
], accent=BLUE)
card(s, L, Y_BODY + 2.25, COL_2, 2.05, "Unit-Tests  ·  White Box", [
    "**xUnit** direkt gegen den `CalculatorService`, ohne Webserver. Laufzeit: "
    "**46 ms für 33 Tests**.",
    "Prüft: Rechenergebnisse, Assoziativität, Ausnahmefälle – also die "
    "**Fachlogik**.",
], accent=AQUA)
bullets(s, L + COL_2 + GAP, Y_BODY, COL_2, [
    "**Testpyramide.** Viele schnelle Unit-Tests unten, wenige aussagekräftige "
    "End-to-End-Tests oben. In der CI laufen die Unit-Tests zuerst und dienen "
    "als schnelles Quality Gate.",
    "**Produktionsnahe Umgebung.** Die API wird in der CI nicht per "
    "`dotnet run` gestartet, sondern als Docker-Container mit "
    "**Produktionskonfiguration** – getestet wird das Artefakt, das ausgeliefert "
    "würde.",
    "**Kein „grün ohne Tests“.** Ist die API nicht erreichbar, schlagen die Tests "
    "in der CI (`CI=true`) hart fehl; lokal werden sie als Inconclusive markiert.",
], size=11.5, gap=16)

# 15 — Testentwurfsverfahren
s = slide("""
Hier zeigt sich Methodik statt Bauchgefühl. Ein Beispiel pro Verfahren:

- Äquivalenzklassen: alle Eingaben mit >= 2 Zahlen verhalten sich gleich – also
  braucht man nicht 50 davon.
- Grenzwertanalyse: die Fehler sitzen an den Rändern. Deshalb 1 / 2 Zahlen und
  1000 / 1001 Zahlen.
- Error Guessing: 0.1 + 0.2 ist in IEEE-754 nicht exakt 0.3. Klassische
  Stolperfalle, also explizit ein Testfall.

Priorisierung: nicht jeder Test ist gleich wichtig. Die Prios stehen als
NUnit-Kategorie im Code, sind also gezielt ausführbar – nicht nur ein Vermerk im
Dokument.
""")
head(s, "Testkonzept", "Testfälle systematisch abgeleitet – nicht geraten",
     foot="Qualitätssicherung", pageno=15)
tf = box(s, L, Y_BODY - 0.03, CW, 0.24)
run(para(tf, first=True), "Black-Box-Testentwurfsverfahren nach ISTQB",
    size=9, color=MUTED)
table(s, L, Y_BODY + 0.26, CW, [0.26, 0.74],
      header=["Verfahren", "Anwendung im Projekt"],
      rows=[
        ["**Äquivalenzklassen**",
         "Gültig (≥ 2 Zahlen; positiv / negativ / dezimal) vs. ungültig (< 2, "
         "> 1000, `null`, falscher Typ, ungültiges JSON)"],
        ["**Grenzwertanalyse**",
         "0 / 1 / 2 Zahlen · 1000 / 1001 Zahlen · Divisor 0 · "
         "Zahlbereichsgrenzen `1e308`, `1e-308`"],
        ["**Vertragsprüfung**",
         "Response-Felder, Statuscodes, Content-Type, unbekannte Routen, falsche "
         "HTTP-Methode"],
        ["**Error Guessing**",
         "Gleitkomma-Präzision (`0.1 + 0.2`), Überlauf-Serialisierung, "
         "unbekannte Zusatzfelder im Body"],
      ], row_h=0.52, size=11)
cy = Y_BODY + 2.72
card(s, L, cy, COL_2, 1.55, "Risikobasierte Priorisierung", [
    "**Prio Hoch** – Happy Path, Division durch null, Eingabevalidierung: "
    "Kernfunktion, betrifft alle Nutzer.",
    "**Prio Mittel** – Gleitkomma-Randfälle, API-Vertrag.",
])
x2 = L + COL_2 + GAP
card(s, x2, cy, COL_2, 1.55, "Im Code hinterlegt",
     "Als NUnit-Kategorie und Allure-Severity, dadurch gezielt ausführbar:")
code(s, x2 + 0.22, cy + 0.98, COL_2 - 0.44,
     ['dotnet test --filter "TestCategory=Prio-Hoch"'], size=9.5)

# 16 — Testumfang
s = slide("""
Kurz auf die Verteilung zeigen. Die Aussage steht in der Karte unten:

„Weniger als die Hälfte der Testfälle prüft den Erfolgsfall. Der größere Teil
prüft Grenzen, Fehler und Vertragsbruch – und genau dort entstehen in der Praxis
die Probleme.“

Bei Nachfrage: Die 9 Validierungsfälle laufen gegen alle vier Endpunkte, die Zahl
ausgeführter Tests ist also deutlich höher als 36.
""")
head(s, "Testumfang", "36 spezifizierte API-Testfälle",
     foot="Qualitätssicherung", pageno=16)
BAR_X, BAR_W, ROW_H = L + 2.75, 6.25, 0.36
for i, (name, val) in enumerate([
    ("Happy Path (4.1)", 14),
    ("Eingabevalidierung (4.4)", 9),
    ("Gleitkomma-Randfälle (4.2)", 6),
    ("API-Vertrag / Robustheit (4.5)", 5),
    ("Division durch null (4.3)", 2),
]):
    y = Y_BODY + i * ROW_H
    tf = box(s, L, y + 0.03, 2.65, 0.26, wrap=False)
    run(para(tf, first=True), name, size=11, color=INK_2)
    rect(s, BAR_X, y + 0.015, BAR_W * val / 14, 0.245, fill=BLUE)
    tf = box(s, BAR_X + BAR_W + 0.12, y + 0.03, 0.5, 0.26)
    run(para(tf, first=True), str(val), size=11, color=INK, bold=True)
hline(s, L, Y_BODY + 5 * ROW_H + 0.06, CW)
tf = box(s, L, Y_BODY + 5 * ROW_H + 0.18, CW, 0.24)
rich(para(tf, first=True),
     "Testfall-IDs je Gruppe im Testkonzept · die Validierungsfälle werden "
     "zusätzlich gegen **alle vier Endpunkte** ausgeführt", size=9.5, color=MUTED)

cy = Y_BODY + 2.32
for i, (val, lbl, note) in enumerate([
    ("33", "Unit-Tests (xUnit)", "46 ms Laufzeit, alle grün"),
    ("36", "Spezifizierte API-Testfälle", "TC-IDs im Testkonzept"),
    ("5", "Getestete Endpunkte", "4 Operationen + Health"),
    ("0", "Offene Defects", "Stand: aktueller Lauf"),
]):
    kw = (CW - 3 * 0.22) / 4
    kx = L + i * (kw + 0.22)
    rect(s, kx, cy, kw, 1.42, fill=PLANE, line=HAIRLINE,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = box(s, kx + 0.22, cy + 0.16, kw - 0.44, 0.55)
    run(para(tf, first=True, line=1.0), val, size=32, color=INK, bold=True)
    tf = box(s, kx + 0.22, cy + 0.80, kw - 0.44, 0.5)
    run(para(tf, first=True, line=1.2), lbl, size=10.5, color=INK_2)
    run(para(tf, space_before=3, line=1.2), note, size=9, color=MUTED)
card(s, L, cy + 1.62, CW, 0.72, None,
     "Auffällig ist die Verteilung: **weniger als die Hälfte** der Testfälle "
     "prüft den Erfolgsfall. Der größere Teil prüft Grenzen, Fehler und "
     "Vertragsbruch – dort entstehen in der Praxis die Probleme.", body_size=11.5)

# 17 — Unit-Tests
s = slide("""
Code zeigen, drei Punkte:
1. Deutsche, sprechende Testnamen – der Bericht liest sich wie eine Spezifikation.
2. [Theory] + [InlineData] – ein neuer Testfall ist eine neue Zeile.
3. [MemberData] prüft eine Regel gegen alle vier Operationen gleichzeitig.
   Verhindert, dass Validierung bei einer Operation vergessen wird.

Und die Toleranz bei double: Assert.Equal(0.3, …, 1e-10). Nie exakt vergleichen.
""")
head(s, "Testautomatisierung", "Unit-Tests: datengetrieben statt kopiert",
     foot="Testautomatisierung", pageno=17)
code(s, L, Y_BODY + 0.28, 6.6, [
    '[Theory]',
    '[InlineData(new double[] { 10, 4 }, 2.5)]',
    '[InlineData(new double[] { 100, 5, 2 }, 10)]',
    '[InlineData(new double[] { -9, 3 }, -3)]',
    '[InlineData(new double[] { 0, 5 }, 0)]',
    'public void Divide_rechnet_verkettet(',
    '        double[] numbers, double expected)',
    '    => Assert.Equal(expected, _service.Divide(numbers));',
    '',
    '// Gleitkomma niemals exakt vergleichen',
    '[Fact]',
    'public void Add_beruecksichtigt_Gleitkomma_Toleranz()',
    '    => Assert.Equal(0.3, _service.Add([0.1, 0.2]), 1e-10);',
    '',
    '// Dieselbe Regel für alle vier Operationen',
    '[Theory]',
    '[MemberData(nameof(Operationen))]',
    'public void Operation_wirft_bei_leerem_Array(string op)',
    '    => Assert.Throws<ArgumentException>(',
    '           () => Invoke(op, []));',
], size=9.5, label="Calculator.Api.Tests/CalculatorServiceTests.cs")
x2 = L + 6.9
bullets(s, x2, Y_BODY + 0.28, CW - 6.9, [
    "**Sprechende deutsche Testnamen** – der Testbericht liest sich als "
    "Spezifikation der Fachlogik.",
    "**`[Theory]` + `[InlineData]`**: ein Test, viele Datensätze. Ein neuer Fall "
    "ist eine neue Zeile.",
    "**`[MemberData]`** prüft eine Regel gleichzeitig gegen alle vier "
    "Operationen – Validierung darf nirgends vergessen werden.",
    "**Toleranzvergleich bei `double`.** `0.1 + 0.2 == 0.3` ist in IEEE-754 "
    "falsch. Der Test prüft gegen `1e-10`.",
], size=11.5, gap=15)
sp = rect(s, x2, Y_BODY + 3.35, CW - 6.9, 0.62, fill=PLANE, line=HAIRLINE,
          shape=MSO_SHAPE.ROUNDED_RECTANGLE)
rect(s, x2 + 0.22, Y_BODY + 3.58, 0.13, 0.13, fill=GOOD,
     shape=MSO_SHAPE.OVAL)
tf = box(s, x2 + 0.44, Y_BODY + 3.51, CW - 7.4, 0.3)
p = para(tf, first=True)
run(p, "Bestanden", size=11, color=GOOD, bold=True)
run(p, " — 33 von 33 Tests, 0 Fehler.", size=11, color=INK_2)

# 18 — API-Tests
s = slide("""
Playwright als HTTP-Client erklären. Dann der Punkt, der methodisch am meisten
wert ist – Traceability:

„Jeder Test trägt die Testfall-ID aus dem Konzept im Namen. Ich kann also von
einem roten Test im Bericht rückwärts zum Testkonzept gehen – und umgekehrt für
jeden spezifizierten Fall nachweisen, dass er automatisiert ist.“

TC-CON-05 als Beispiel erwähnen: unbekannte Zusatzfelder werden toleriert. Das
ist eine dokumentierte Entscheidung, kein Zufall – und deshalb ein Testfall.
""")
head(s, "Testautomatisierung", "API-Tests: Playwright ohne Browser",
     foot="Testautomatisierung", pageno=18)
bullets(s, L, Y_BODY, COL_2, [
    "**`APIRequestContext`** statt Browser: Playwright wird hier ausschließlich "
    "als HTTP-Client genutzt – schnell, stabil, kein Rendering.",
    "**Parallelisierung auf Fixture-Ebene** (NUnit `Parallelizable`); jede "
    "Fixture bekommt ihre eigene Playwright-Instanz. Möglich, weil die API "
    "zustandslos ist.",
    "**Konfigurierbare Basis-URL** über `API_BASE_URL` – lokal "
    "`localhost:5116`, in der CI `http://calculator-api:8080` im Docker-Netzwerk.",
    "**Entry-Kriterium im Code:** vor dem Lauf wird `GET /health` geprüft.",
], size=11.5, gap=16)
x2 = L + COL_2 + GAP
card(s, x2, Y_BODY, COL_2, 2.45, "Rückverfolgbarkeit (Traceability)",
     "Jeder automatisierte Test trägt seine **Testfall-ID** aus dem Testkonzept "
     "im Namen:")
code(s, x2 + 0.22, Y_BODY + 0.95, COL_2 - 0.44,
     ['TC-ADD-01 add 1+2=3', 'TC-DIV0-01 divide [10, 0] → 400'], size=9.5)
tf = box(s, x2 + 0.22, Y_BODY + 1.78, COL_2 - 0.44, 0.55)
rich(para(tf, first=True, line=1.3),
     "Damit ist die Kette **Konzept ↔ Code ↔ Report** lückenlos "
     "nachvollziehbar – in beide Richtungen.", size=10.5)
card(s, x2, Y_BODY + 2.68, COL_2, 1.35, "Beispiel eines Robustheits-Testfalls",
     "`TC-CON-05`: Unbekannte Zusatzfelder im Body werden **toleriert** (200) – "
     "bewusste Entscheidung nach dem Robustheitsprinzip, nicht ein Versehen.")

# 19 — Docker
s = slide("""
Multi-Stage in einem Satz: gebaut mit dem SDK-Image, ausgeliefert auf dem
schlanken Runtime-Image.

Die Layer-Reihenfolge erklären, wenn Zeit ist: erst csproj kopieren, dann
restore, dann der Rest. Solange sich Abhängigkeiten nicht ändern, kommt der
Restore aus dem Cache – spart bei jedem Build Zeit.
""")
head(s, "Betrieb", "Containerisierung: Multi-Stage-Build",
     foot="Betrieb", pageno=19)
code(s, L, Y_BODY + 0.28, 6.6, [
    '# Build-Stage',
    'FROM mcr.microsoft.com/dotnet/sdk:10.0 AS build',
    'COPY …/Calculator.Api.csproj Calculator.Api/',
    'RUN dotnet restore Calculator.Api/Calculator.Api.csproj',
    'COPY Calculator.Api/Calculator.Api/ Calculator.Api/',
    'RUN dotnet publish … -c Release -o /app',
    '',
    '# Runtime-Stage',
    'FROM mcr.microsoft.com/dotnet/aspnet:10.0 AS runtime',
    'COPY --from=build /app .',
    'ENV ASPNETCORE_HTTP_PORTS=8080',
    'EXPOSE 8080',
    '',
    '# Nicht als root laufen',
    'USER app',
    '',
    'HEALTHCHECK --interval=30s --retries=3 \\',
    '  CMD curl -fsS http://localhost:8080/health || exit 1',
    '',
    'ENTRYPOINT ["dotnet", "Calculator.Api.dll"]',
], size=9.5, label="Dockerfile")
x2 = L + 6.9
bullets(s, x2, Y_BODY + 0.28, CW - 6.9, [
    "**Zwei Stages.** Gebaut wird mit dem SDK-Image, ausgeliefert wird auf dem "
    "schlanken Runtime-Image – das SDK landet nicht im Ergebnis.",
    "**`csproj` zuerst kopieren**, dann `restore`, dann der restliche Code: der "
    "Restore-Layer wird aus dem Docker-Cache bedient, solange sich die "
    "Abhängigkeiten nicht ändern.",
    "**`USER app`** – der Prozess läuft nicht als root. Kleine Zeile, große "
    "Wirkung für die Angriffsfläche.",
    "**`HEALTHCHECK`** macht den Zustand des Containers von außen sichtbar – "
    "dieselbe Probe, die auch CI und Tests verwenden.",
], size=11.5, gap=15)

# 20 — Pipeline
s = slide("""
Die sechs Schritte durchgehen. Zwei Details hervorheben, die Erfahrung zeigen:

- continue-on-error beim Testschritt: absichtlich. Sonst gibt es bei roten Tests
  keinen Report – und genau dann braucht man ihn. Am Ende schlägt der Workflow
  explizit fehl.
- docker logs bei Fehlern: die API-Logs landen automatisch im Lauf. Man muss
  nicht raten, warum ein Test rot war.
""")
head(s, "CI/CD", "Die Pipeline: von Push zu Testbericht", foot="CI/CD", pageno=20)
steps = [
    "**Unit-Tests** ausführen — schnelles Quality Gate, bricht früh ab",
    "**Test-Repo auschecken** — Ref per `workflow_dispatch` wählbar",
    "**Zwei Docker-Images bauen** — API und Tests",
    "**API starten** im Docker-Netzwerk, auf `/health` warten (max. 60 s)",
    "**Tests im Container** gegen die laufende API ausführen",
    "**Ergebnisse veröffentlichen** — Artefakt, TRX-Check, Allure-Report",
]
for i, text in enumerate(steps):
    y = Y_BODY + i * 0.72
    rect(s, L, y, COL_2, 0.60, fill=PLANE, line=HAIRLINE,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    n = rect(s, L + 0.17, y + 0.15, 0.30, 0.30, fill=BLUE,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    ntf = n.text_frame
    ntf.vertical_anchor = MSO_ANCHOR.MIDDLE
    ntf.margin_left = ntf.margin_right = 0
    ntf.margin_top = ntf.margin_bottom = 0
    run(para(ntf, first=True, align=PP_ALIGN.CENTER), str(i + 1),
        size=10.5, color=WHITE, bold=True)
    tf = box(s, L + 0.62, y + 0.10, COL_2 - 0.82, 0.42,
             anchor=MSO_ANCHOR.MIDDLE)
    rich(para(tf, first=True, line=1.2), text, size=10.5)
bullets(s, L + COL_2 + GAP, Y_BODY, COL_2, [
    "**Trigger:** Push und Pull Request auf `main`, zusätzlich manuell.",
    "**`continue-on-error` mit Absicht.** Der Test-Schritt bricht nicht sofort "
    "ab – so werden Report und Artefakte auch bei roten Tests erzeugt. Am Ende "
    "schlägt der Workflow dann explizit fehl.",
    "**Diagnose eingebaut:** bei Testfehlern werden automatisch die "
    "`docker logs` der API ausgegeben.",
    "**PR-Gate:** Ein Pull Request kann nur gemerged werden, wenn die Pipeline "
    "grün ist. Für PRs wird bewusst nicht auf GitHub Pages veröffentlicht.",
], size=11.5, gap=16)

# 21 — Reporting
s = slide("""
„Ein Testlauf, dessen Ergebnis niemand ansieht, ist verschwendete Rechenzeit.“

Drei Ausgabekanäle: Allure-Report auf GitHub Pages (visuell, mit Historie über
20 Läufe), TRX-Check direkt am Commit, Rohdaten als Artefakt.

Wenn Internet verfügbar ist: hier den Report live öffnen. Das ist der stärkste
Moment der Präsentation.
""")
head(s, "CI/CD", "Reporting: Ergebnisse, die man ansehen kann",
     foot="CI/CD", pageno=21)
for i, (t, b) in enumerate([
    ("Allure-Report",
     "Visueller Bericht mit Testfall-IDs, Severity und Fehlerdetails – "
     "veröffentlicht auf **GitHub Pages**, versioniert pro Lauf-Nummer."),
    ("Historie",
     "Die **letzten 20 Läufe** bleiben erhalten. Trends und instabile Tests "
     "werden dadurch über die Zeit sichtbar, nicht nur der aktuelle Stand."),
    ("TRX-Check & Artefakt",
     "Ergebnisse erscheinen als **Check direkt am Commit**; die Rohdaten liegen "
     "14 Tage als Workflow-Artefakt bereit."),
]):
    card(s, L + i * (COL_3 + 0.24), Y_BODY, COL_3, 1.95, t, b)
rect(s, L, Y_BODY + 2.20, CW, 1.55, fill=PLANE, line=HAIRLINE,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE)
tf = box(s, L + 0.28, Y_BODY + 2.42, CW - 0.56, 0.28)
run(para(tf, first=True), "Live abrufbar:", size=11.5, color=INK, bold=True)
tf = box(s, L + 0.28, Y_BODY + 2.76, CW - 0.56, 0.32)
run(para(tf, first=True), "https://sascha-pommernell.github.io/calculator-api/",
    size=14, color=BLUE, font=MONO)
tf = box(s, L + 0.28, Y_BODY + 3.18, CW - 0.56, 0.4)
rich(para(tf, first=True, line=1.3),
     "Der Link zum jeweiligen Lauf steht zusätzlich in der Zusammenfassung des "
     "Actions-Laufs – niemand muss nach dem Bericht suchen.", size=11)

# 22 — Herausforderungen
s = slide("""
Ehrlich sein – diese Folie macht die Präsentation glaubwürdig. Die erste Zeile ist
die wichtigste:

„Zwei Testfälle zum Überlauf sind fehlgeschlagen. Das war kein Testfehler,
sondern ein echter Defect: die API hat Infinity ausgeliefert. Der Test hat also
nicht nur geprüft, sondern das Design verbessert – genau dafür schreibt man ihn.“

Die letzte Zeile als Abschluss:

„Und die unangenehmste Erkenntnis: meine Pipeline war einmal grün, ohne dass
Tests gelaufen sind. Ein grüner Haken muss etwas bedeuten, sonst ist er
schädlich.“
""")
head(s, "Erfahrungen", "Herausforderungen & Lösungen",
     foot="Erfahrungen", pageno=22)
table(s, L, Y_BODY, CW, [0.34, 0.66],
      header=["Problem", "Lösung"],
      rows=[
        ["**Überlauf lieferte `Infinity` statt eines Fehlers**",
         "Zwei Testfälle (TC-FLT-02/03) schlugen fehl und deckten einen echten "
         "Defect auf. Ergebnis: die `double.IsFinite`-Prüfung im Service – "
         "**der Test hat das Design verbessert**."],
        ["**Testergebnisse aus dem Container nicht lesbar**",
         "Der Container schrieb als root in das gemountete Volume; der Runner "
         "konnte die Dateien nicht weiterverarbeiten. Gelöst durch expliziten "
         "`chown`-Schritt nach dem Testlauf."],
        ["**Allure-Action lieferte keinen Report mehr**",
         "Statt einer fremden Action wird die **Allure CLI in fixierter Version** "
         "installiert und der Report direkt generiert – weniger Abhängigkeiten, "
         "reproduzierbar."],
        ["**Erster Lauf brach ab: `gh-pages` existierte nicht**",
         "Der Workflow prüft jetzt per `git ls-remote`, ob der Branch existiert, "
         "und überspringt das Laden der Historie beim ersten Lauf."],
        ["**Pipeline war grün, obwohl keine Tests liefen**",
         "Bei nicht erreichbarer API wurde übersprungen statt gemeldet. Jetzt: "
         "`CI=true` erzwingt harten Fehlschlag – **ein grüner Haken muss etwas "
         "bedeuten**."],
      ], row_h=0.90, size=11)

# 23 — Fazit
s = slide("""
Sechs Punkte, nicht einzeln vorlesen. Die Kernaussage aus der Karte sprechen:

„Das Ergebnis dieses Projekts ist nicht die Rechenlogik. Es ist der belegbare,
wiederholbare Weg von der Anforderung zum getesteten, ausgelieferten Artefakt.“
""")
head(s, "Fazit", "Was das Projekt zeigt", foot="Fazit", pageno=23)
bullets(s, L, Y_BODY, COL_2, [
    "**Saubere Architektur** — Trennung von HTTP-Schicht, Fachlogik und "
    "Fehlerbehandlung; Abhängigkeiten über Interfaces.",
    "**Der Vertrag steht im Mittelpunkt** — einheitliche ProblemDetails, "
    "OpenAPI, geprüft durch Vertragstests.",
    "**Systematischer Testentwurf** — Äquivalenzklassen, Grenzwerte, "
    "Priorisierung nach Risiko; keine zufällig entstandenen Tests.",
], size=12, gap=18)
x2 = L + COL_2 + GAP
bullets(s, x2, Y_BODY, COL_2, [
    "**Zwei Testebenen** — schnelle Unit-Tests als Gate, Black-Box-API-Tests "
    "gegen das echte Artefakt.",
    "**Durchgängige Automatisierung** — Build, Container, Test und "
    "Veröffentlichung bei jedem Push, ohne Handarbeit.",
    "**Nachvollziehbarkeit** — jede Zeile im Testbericht ist bis zum "
    "Testkonzept zurückverfolgbar.",
], size=12, gap=18)
card(s, x2, Y_BODY + 2.85, COL_2, 1.15, None,
     "**Kernaussage:** Nicht die Rechenlogik ist das Ergebnis – sondern der "
     "belegbare, wiederholbare Weg von der Anforderung zum getesteten, "
     "ausgelieferten Artefakt.", body_size=11.5)

# 24 — Ausblick
s = slide("""
Zeigt, dass die Grenzen des Projekts bekannt sind. Besonders erwähnen:
Rate Limiting ist implementiert, aber nicht unter Last verifiziert – eine bewusst
offene Lücke, nicht ein Versehen.
""")
head(s, "Ausblick", "Nächste sinnvolle Schritte", foot="Ausblick", pageno=24)
for i, (t, b) in enumerate([
    ("Authentifizierung",
     "Bewusst nicht im Scope. Mit API-Keys oder JWT käme eine ganze Testklasse "
     "hinzu: 401/403-Fälle, Token-Ablauf, Rollen."),
    ("Last- und Performancetests",
     "Das Rate Limit ist implementiert, aber nicht unter Last verifiziert. Ein "
     "k6- oder NBomber-Lauf in der Pipeline würde das messen."),
    ("Code-Coverage-Gate",
     "Coverage messen und als Schwellwert in der Pipeline erzwingen, damit neue "
     "Logik nicht ungetestet einzieht."),
    ("Weitere Operationen",
     "Potenz, Modulo, Wurzel – jeweils nach derselben Regel: Testfälle "
     "(Happy Path + Negativ) **vor** dem Merge."),
    ("Vertragstests gegen OpenAPI",
     "Responses automatisch gegen das OpenAPI-Schema validieren, statt Felder "
     "einzeln zu prüfen."),
    ("Strukturiertes Logging",
     "Korrelations-IDs und strukturierte Logs, um Requests in einem "
     "Produktionsbetrieb nachverfolgen zu können."),
]):
    card(s, L + (i % 3) * (COL_3 + 0.24), Y_BODY + (i // 3) * 2.28,
         COL_3, 2.05, t, b)

# 25 — Abschluss
s = slide("""
Danken, zu Fragen einladen. Swagger UI und Allure-Report bereithalten.

Vorbereitete Antworten:

„Warum double und nicht decimal?“ – double passt zum JSON-Zahlentyp. decimal wäre
für Geldbeträge richtig. Die Konsequenzen (Präzisionstoleranz, Überlaufprüfung)
sind explizit behandelt und getestet.

„Warum Exceptions statt Result-Objekte?“ – Exceptions können nicht versehentlich
ignoriert werden; der zentrale IExceptionHandler übersetzt sie an einer Stelle.

„Ist Playwright für API-Tests überdimensioniert?“ – Es bringt APIRequestContext,
Parallelisierung und Tracing mit, und käme eine UI hinzu, wäre kein
Werkzeugwechsel nötig.

„Wie hoch ist die Code-Coverage?“ – Wird nicht gemessen, steht im Ausblick. Die
Fachlogik ist durch 33 Unit-Tests inkl. aller Ausnahmepfade abgedeckt.

„Warum 200 Requests pro Sekunde?“ – Pragmatisch: hoch genug für parallele
Testläufe, niedrig genug um eine Schleife zu bremsen. Ohne Lasttest wäre jede
exaktere Zahl geraten.

„Wie lange läuft die Pipeline?“ – Dominiert von den zwei Docker-Builds; die
Unit-Tests sind in Millisekunden fertig und laufen deshalb als erstes.
""")
tf = box(s, L, 2.75, 8, 0.3)
run(para(tf, first=True), "ENDE", size=11, color=BLUE, bold=True)
tf = box(s, L, 3.15, 10, 0.85)
run(para(tf, first=True, line=1.05), "Vielen Dank.", size=40, color=INK, bold=True)
tf = box(s, L, 4.25, 8.2, 0.5)
rich(para(tf, first=True, line=1.32),
     "Fragen? Gerne auch direkt an der laufenden API und am Allure-Report.",
     size=15)
tf = box(s, L, 5.0, 8.2, 0.7)
run(para(tf, first=True, line=1.5),
    "github.com/Sascha-Pommernell/calculator-api", size=11.5, color=BLUE, font=MONO)
run(para(tf, line=1.5), "sascha-pommernell.github.io/calculator-api",
    size=11.5, color=BLUE, font=MONO)


out = Path(__file__).resolve().parent / "Calculator-API-Praesentation.pptx"
prs.save(out)
print(f"{out}  ({len(prs.slides.__iter__.__self__._sldIdLst)} Folien)")
